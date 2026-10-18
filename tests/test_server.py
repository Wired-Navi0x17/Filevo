"""Step 2 acceptance tests -- the FastAPI backend behind the Sortwise UI.

The tests exercise the real service layer (extraction, safe move, collision
handling, undo ledger) against an injected *fake* classifier, so the suite runs
in milliseconds without downloading or loading ``all-MiniLM-L6-v2``.

Covered endpoints: ``/api/config``, ``/api/subjects``, ``/api/scan``,
``/api/route`` (single + batch), ``/api/rename/suggest``, ``/api/history`` and
``/api/history/{id}/undo``.
"""

from __future__ import annotations

import sys
from pathlib import Path

import pytest

PROJECT_ROOT = Path(__file__).resolve().parents[1]
if str(PROJECT_ROOT) not in sys.path:
    sys.path.insert(0, str(PROJECT_ROOT))

import docx  # python-docx
import fitz  # PyMuPDF
from fastapi.testclient import TestClient
from pptx import Presentation

from src.calibrator import CalibratedDecision, STATUS_CONFIDENT
from src.server import (
    create_app,
    human_size,
    sanitize_filename,
    unique_path,
)
from src.taxonomy import subject_labels

AI_SUBJECT = "CS2227 - AI and ML"
ALGORITHMS_SUBJECT = "CS2000 - Design and Analysis of Algorithms"


# --------------------------------------------------------------------------- #
# Test doubles
# --------------------------------------------------------------------------- #
class FakeClassifier:
    """Stand-in for :class:`~src.calibrator.CalibratedClassifier`."""

    def __init__(
        self,
        subject: str = AI_SUBJECT,
        confidence: float = 0.93,
        entropy: float = 0.42,
    ) -> None:
        self.subject = subject
        self.confidence = confidence
        self.entropy = entropy
        self._model = object()  # marks the model as loaded for /api/health

    def classify_document(self, doc) -> CalibratedDecision:  # noqa: ANN001
        others = [label for label in subject_labels() if label != self.subject]
        share = max(0.0, 1.0 - self.confidence) / max(1, len(others))
        ranking = [(self.subject, self.confidence)]
        ranking += [(label, share) for label in others]
        return CalibratedDecision(
            subject=self.subject,
            confidence=self.confidence,
            entropy=self.entropy,
            status=STATUS_CONFIDENT,
            ranking=ranking,
        )


# --------------------------------------------------------------------------- #
# Sample document factories
# --------------------------------------------------------------------------- #
def write_docx(path: Path, title: str = "Machine Learning Notes") -> Path:
    document = docx.Document()
    document.core_properties.title = title
    document.add_heading("Supervised Learning", level=1)
    document.add_paragraph(
        "Gradient descent and backpropagation train neural networks on a "
        "training set, with validation and a bias-variance tradeoff."
    )
    document.save(str(path))
    return path


def write_pdf(path: Path, title: str = "Algorithms Lecture Notes") -> Path:
    doc = fitz.open()
    page = doc.new_page(width=595, height=842)
    page.insert_text((72, 60), title, fontsize=20)
    page.insert_textbox(
        fitz.Rect(72, 100, 523, 700),
        "Dynamic programming solves overlapping subproblems. Dijkstra and "
        "Bellman-Ford compute shortest paths; NP-completeness covers "
        "polynomial reductions and approximation algorithms.",
        fontsize=11,
    )
    doc.set_metadata({"title": title})
    doc.save(str(path))
    doc.close()
    return path


def write_pptx(path: Path, title: str = "Cyber Security Slides") -> Path:
    presentation = Presentation()
    slide = presentation.slides.add_slide(presentation.slide_layouts[0])
    slide.shapes.title.text = "Cryptography Fundamentals"
    slide.placeholders[1].text = "AES, RSA, hash functions and Diffie-Hellman key exchange."
    presentation.core_properties.title = title
    presentation.save(str(path))
    return path


# --------------------------------------------------------------------------- #
# Fixtures
# --------------------------------------------------------------------------- #
@pytest.fixture()
def sandbox(tmp_path: Path) -> dict:
    watch = tmp_path / "watch"
    watch.mkdir()
    return {
        "root": tmp_path,
        "watch": watch,
        "out": tmp_path / "out",
        "config": tmp_path / "config.json",
        "ledger": tmp_path / "history.json",
    }


@pytest.fixture()
def docs(sandbox: dict) -> dict[str, Path]:
    watch = sandbox["watch"]
    return {
        "docx": write_docx(watch / "lecture.docx"),
        "pdf": write_pdf(watch / "algorithms.pdf"),
        "pptx": write_pptx(watch / "security.pptx"),
    }


def build_client(sandbox: dict, classifier: FakeClassifier | None = None) -> TestClient:
    app = create_app(
        config_path=sandbox["config"],
        ledger_path=sandbox["ledger"],
        classifier=classifier or FakeClassifier(),
    )
    return TestClient(app)


@pytest.fixture()
def client(sandbox: dict) -> TestClient:
    """Client pointed at the sandbox watch folder and target root."""
    with build_client(sandbox) as test_client:
        test_client.post(
            "/api/config",
            json={
                "watch_folder": str(sandbox["watch"]),
                "target_root": str(sandbox["out"]),
                "auto_move": False,
                "recursive": False,
            },
        )
        yield test_client


# --------------------------------------------------------------------------- #
# /api/config
# --------------------------------------------------------------------------- #
def test_config_defaults_are_well_formed(sandbox: dict) -> None:
    with build_client(sandbox) as test_client:
        response = test_client.get("/api/config")
    assert response.status_code == 200
    config = response.json()
    for key in (
        "watch_folder",
        "target_root",
        "confidence_threshold",
        "entropy_threshold",
        "auto_move",
        "extensions",
        "subject_folder_template",
        "rename_template",
    ):
        assert key in config
    assert 0.0 <= config["confidence_threshold"] <= 1.0
    assert set(config["extensions"]) == {"pdf", "docx", "pptx"}


def test_config_update_persists_to_disk(client: TestClient, sandbox: dict) -> None:
    response = client.post(
        "/api/config",
        json={
            "confidence_threshold": 0.55,
            "entropy_threshold": 2.0,
            "auto_move": True,
        },
    )
    assert response.status_code == 200
    updated = response.json()
    assert updated["confidence_threshold"] == pytest.approx(0.55)
    assert updated["auto_move"] is True

    # The change must survive a fresh GET (i.e. it was written to disk).
    assert client.get("/api/config").json()["confidence_threshold"] == pytest.approx(0.55)
    assert sandbox["config"].exists()


def test_config_rejects_out_of_range_threshold(client: TestClient) -> None:
    response = client.post("/api/config", json={"confidence_threshold": 5.0})
    assert response.status_code == 422


def test_config_empty_body_is_a_no_op(client: TestClient) -> None:
    assert client.put("/api/config").status_code == 200


# --------------------------------------------------------------------------- #
# /api/subjects
# --------------------------------------------------------------------------- #
def test_subjects_exposes_the_full_taxonomy(client: TestClient, sandbox: dict) -> None:
    payload = client.get("/api/subjects").json()
    assert payload["count"] == 8
    labels = [item["label"] for item in payload["subjects"]]
    assert set(labels) == set(subject_labels())

    for item in payload["subjects"]:
        assert item["code"]
        assert item["name"]
        assert item["keywords"]
        # Every subject resolves to a folder under the configured target root.
        assert item["folder"].startswith(str(sandbox["out"]))


def test_subject_folder_template_is_applied(client: TestClient, sandbox: dict) -> None:
    client.post("/api/config", json={"subject_folder_template": "{code}"})
    payload = client.get("/api/subjects").json()
    folders = {item["code"]: item["folder"] for item in payload["subjects"]}
    assert folders["CS2227"] == str(sandbox["out"] / "CS2227")


# --------------------------------------------------------------------------- #
# /api/scan
# --------------------------------------------------------------------------- #
def test_scan_finds_only_supported_documents(client: TestClient, docs: dict, sandbox: dict) -> None:
    (sandbox["watch"] / "notes.txt").write_text("ignore me", encoding="utf-8")
    (sandbox["watch"] / ".hidden.pdf").write_text("x", encoding="utf-8")

    payload = client.post("/api/scan", json={}).json()
    names = {item["name"] for item in payload["files"]}
    assert names == {"lecture.docx", "algorithms.pdf", "security.pptx"}
    assert payload["count"] == 3
    assert payload["folder"] == str(sandbox["watch"])

    sample = payload["files"][0]
    for key in ("name", "path", "relative", "file_type", "size", "size_human", "modified"):
        assert key in sample
    assert sample["file_type"] in {"pdf", "docx", "pptx"}


def test_scan_is_also_available_via_get(client: TestClient, docs: dict) -> None:
    assert client.get("/api/scan").status_code == 200
    assert client.get("/api/scan").json()["count"] == 3


def test_scan_custom_path(client: TestClient, docs: dict, sandbox: dict) -> None:
    payload = client.post("/api/scan", json={"path": str(sandbox["watch"])}).json()
    assert payload["count"] == 3


def test_scan_missing_folder_returns_404(client: TestClient, sandbox: dict) -> None:
    response = client.post("/api/scan", json={"path": str(sandbox["root"] / "nope")})
    assert response.status_code == 404
    assert "does not exist" in response.json()["detail"]


# --------------------------------------------------------------------------- #
# /api/route
# --------------------------------------------------------------------------- #
def test_route_dry_run_does_not_move(client: TestClient, docs: dict, sandbox: dict) -> None:
    response = client.post(
        "/api/route",
        json={"path": str(docs["docx"]), "move": True, "dry_run": True},
    )
    assert response.status_code == 200
    result = response.json()
    assert result["moved"] is False
    assert result["dry_run"] is True
    assert result["subject"] == AI_SUBJECT
    assert result["status"] == "CONFIDENT"
    assert result["confidence"] == pytest.approx(0.93)
    assert result["destination"]["folder"] == str(sandbox["out"] / AI_SUBJECT)
    assert docs["docx"].exists(), "dry run must not touch the source file"
    # Nothing was moved, so the ledger stays empty.
    assert client.get("/api/history").json()["count"] == 0


def test_route_moves_and_records_history(client: TestClient, docs: dict, sandbox: dict) -> None:
    response = client.post(
        "/api/route",
        json={"path": str(docs["docx"]), "move": True},
    )
    assert response.status_code == 200
    result = response.json()

    assert result["moved"] is True
    destination = Path(result["destination"]["path"])
    assert destination.exists()
    assert destination.parent == sandbox["out"] / AI_SUBJECT
    assert not docs["docx"].exists()
    assert result["entry_id"]

    history = client.get("/api/history").json()
    assert history["count"] == 1
    entry = history["entries"][0]
    assert entry["subject"] == AI_SUBJECT
    assert entry["destination"] == str(destination)
    assert entry["undone"] is False


def test_route_respects_auto_move_config(client: TestClient, docs: dict) -> None:
    # Off (default): the file stays put unless the caller asks for the move.
    assert client.post("/api/route", json={"path": str(docs["pdf"])}).json()["moved"] is False
    assert docs["pdf"].exists()

    client.post("/api/config", json={"auto_move": True})
    result = client.post("/api/route", json={"path": str(docs["pdf"])}).json()
    assert result["moved"] is True
    assert not docs["pdf"].exists()


def test_route_renames_in_flight(client: TestClient, docs: dict) -> None:
    result = client.post(
        "/api/route",
        json={"path": str(docs["pptx"]), "rename": "Week 4 - Cryptography", "move": True},
    ).json()
    assert result["destination"]["name"] == "Week 4 - Cryptography.pptx"
    assert result["renamed_from"] == "security.pptx"
    assert result["destination"]["renamed"] is True
    assert Path(result["destination"]["path"]).exists()


def test_route_sanitises_dangerous_rename(client: TestClient, docs: dict) -> None:
    result = client.post(
        "/api/route",
        json={"path": str(docs["pptx"]), "rename": "../../etc/pa ss:wd", "move": False},
    ).json()
    name = result["destination"]["name"]
    assert "/" not in name and "\\" not in name and ":" not in name
    assert name.endswith(".pptx")


def test_route_adds_extension_when_missing(client: TestClient, docs: dict) -> None:
    result = client.post(
        "/api/route",
        json={"path": str(docs["docx"]), "rename": "clean-name", "move": False},
    ).json()
    assert result["destination"]["name"] == "clean-name.docx"


def test_route_collision_gets_a_suffix(client: TestClient, docs: dict, sandbox: dict) -> None:
    occupied = sandbox["out"] / AI_SUBJECT
    occupied.mkdir(parents=True, exist_ok=True)
    (occupied / "lecture.docx").write_bytes(b"already here")

    result = client.post(
        "/api/route", json={"path": str(docs["docx"]), "move": True}
    ).json()
    assert result["destination"]["name"] == "lecture (1).docx"
    assert result["destination"]["collision"] is True
    assert (occupied / "lecture.docx").read_bytes() == b"already here"
    assert (occupied / "lecture (1).docx").exists()


def test_route_manual_subject_override(client: TestClient, docs: dict, sandbox: dict) -> None:
    result = client.post(
        "/api/route",
        json={"path": str(docs["docx"]), "subject": ALGORITHMS_SUBJECT, "move": True},
    ).json()
    assert result["subject"] == ALGORITHMS_SUBJECT
    assert result["manual"] is True
    assert Path(result["destination"]["path"]).parent == sandbox["out"] / ALGORITHMS_SUBJECT


def test_route_unknown_subject_is_rejected(client: TestClient, docs: dict) -> None:
    response = client.post(
        "/api/route", json={"path": str(docs["docx"]), "subject": "CS9999 - Nope"}
    )
    assert response.status_code == 400


def test_route_unsupported_type_returns_415(client: TestClient, sandbox: dict) -> None:
    bogus = sandbox["watch"] / "notes.txt"
    bogus.write_text("plain text", encoding="utf-8")
    response = client.post("/api/route", json={"path": str(bogus)})
    assert response.status_code == 415


def test_route_missing_file_returns_404(client: TestClient, sandbox: dict) -> None:
    response = client.post("/api/route", json={"path": str(sandbox["watch"] / "ghost.pdf")})
    assert response.status_code == 404


def test_route_corrupt_document_returns_422(client: TestClient, sandbox: dict) -> None:
    corrupt = sandbox["watch"] / "broken.pdf"
    corrupt.write_bytes(b"%PDF-1.4 this is not really a pdf")
    response = client.post("/api/route", json={"path": str(corrupt)})
    assert response.status_code == 422


def test_route_requires_a_path(client: TestClient) -> None:
    assert client.post("/api/route", json={}).status_code == 400


def test_status_follows_configured_threshold(sandbox: dict, docs: dict) -> None:
    strict = FakeClassifier(confidence=0.60, entropy=0.9)
    with build_client(sandbox, strict) as test_client:
        test_client.post(
            "/api/config",
            json={
                "watch_folder": str(sandbox["watch"]),
                "target_root": str(sandbox["out"]),
                "confidence_threshold": 0.80,
            },
        )
        result = test_client.post("/api/route", json={"path": str(docs["docx"])}).json()
    assert result["status"] == "AMBIGUOUS"
    assert result["confidence"] == pytest.approx(0.60)


def test_route_batch(sandbox: dict, docs: dict) -> None:
    with build_client(sandbox) as test_client:
        test_client.post(
            "/api/config",
            json={"watch_folder": str(sandbox["watch"]), "target_root": str(sandbox["out"])},
        )
        response = test_client.post(
            "/api/route",
            json={"paths": [str(docs["docx"]), str(docs["pdf"])], "move": True},
        )
        assert response.status_code == 200
        payload = response.json()
        assert payload["count"] == 2
        assert payload["moved"] == 2

        dedicated = test_client.post(
            "/api/route/batch",
            json={"paths": [str(docs["pptx"])], "move": True},
        ).json()
        assert dedicated["moved"] == 1


# --------------------------------------------------------------------------- #
# /api/rename/suggest
# --------------------------------------------------------------------------- #
def test_rename_suggestion_uses_document_title(client: TestClient, docs: dict) -> None:
    payload = client.post("/api/rename/suggest", json={"path": str(docs["docx"])}).json()
    assert payload["title"] == "Machine Learning Notes"
    assert payload["suggested"] == "Machine Learning Notes.docx"


def test_rename_suggestion_honours_template(client: TestClient, docs: dict) -> None:
    client.post("/api/config", json={"rename_template": "{code} - {title}"})
    payload = client.post(
        "/api/rename/suggest",
        json={"path": str(docs["docx"]), "subject": AI_SUBJECT},
    ).json()
    assert payload["suggested"] == "CS2227 - Machine Learning Notes.docx"


# --------------------------------------------------------------------------- #
# history + undo
# --------------------------------------------------------------------------- #
def test_undo_restores_the_file(client: TestClient, docs: dict, sandbox: dict) -> None:
    original = docs["docx"]
    routed = client.post("/api/route", json={"path": str(original), "move": True}).json()
    assert not original.exists()

    response = client.post(f"/api/history/{routed['entry_id']}/undo")
    assert response.status_code == 200
    body = response.json()
    assert Path(body["restored_to"]) == original
    assert original.exists(), "undo must put the file back where it came from"

    history = client.get("/api/history").json()
    entries = {entry["id"]: entry for entry in history["entries"]}
    assert entries[routed["entry_id"]]["undone"] is True
    assert any(entry["action"] == "undo" for entry in history["entries"])


def test_undo_twice_is_a_conflict(client: TestClient, docs: dict) -> None:
    routed = client.post("/api/route", json={"path": str(docs["pdf"]), "move": True}).json()
    assert client.post(f"/api/history/{routed['entry_id']}/undo").status_code == 200
    assert client.post(f"/api/history/{routed['entry_id']}/undo").status_code == 409


def test_undo_unknown_entry_returns_404(client: TestClient) -> None:
    assert client.post("/api/history/deadbeef/undo").status_code == 404


def test_history_can_be_cleared(client: TestClient, docs: dict) -> None:
    client.post("/api/route", json={"path": str(docs["pptx"]), "move": True})
    assert client.get("/api/history").json()["count"] == 1
    assert client.delete("/api/history").json()["cleared"] == 1
    assert client.get("/api/history").json()["count"] == 0


# --------------------------------------------------------------------------- #
# health + static assets
# --------------------------------------------------------------------------- #
def test_health_reports_telemetry(client: TestClient, docs: dict) -> None:
    payload = client.get("/api/health").json()
    assert payload["status"] == "online"
    assert payload["model_loaded"] is True
    assert payload["queue_count"] == 3
    assert payload["uptime_seconds"] >= 0


def test_index_and_assets_are_served(client: TestClient) -> None:
    index = client.get("/")
    assert index.status_code == 200
    assert "SORTWISE" in index.text

    for asset in ("/static/styles.css", "/static/app.js", "/static/sortwise.svg"):
        assert client.get(asset).status_code == 200, asset
    assert client.get("/favicon.svg").status_code == 200


# --------------------------------------------------------------------------- #
# pure helpers
# --------------------------------------------------------------------------- #
@pytest.mark.parametrize(
    "raw,expected",
    [
        ("Normal Name", "Normal Name"),
        ("a/b\\c:d", "a-b-c-d"),
        ("  spaced  out  ", "spaced out"),
        ("trailing...", "trailing"),
        ("", "fallback"),
        ("<>|?*", "-----"),
    ],
)
def test_sanitize_filename(raw: str, expected: str) -> None:
    fallback = "fallback"
    assert sanitize_filename(raw, fallback) == expected


def test_unique_path_avoids_collisions(tmp_path: Path) -> None:
    (tmp_path / "notes.pdf").write_text("a", encoding="utf-8")
    (tmp_path / "notes (1).pdf").write_text("b", encoding="utf-8")
    assert unique_path(tmp_path, "notes.pdf").name == "notes (2).pdf"
    assert unique_path(tmp_path, "fresh.pdf").name == "fresh.pdf"


@pytest.mark.parametrize(
    "size,expected",
    [(0, "0 B"), (512, "512 B"), (2048, "2.0 KB"), (5 * 1024 * 1024, "5.0 MB"), (None, "-")],
)
def test_human_size(size, expected) -> None:
    assert human_size(size) == expected
