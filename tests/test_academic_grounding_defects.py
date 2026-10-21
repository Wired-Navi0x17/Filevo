import os
import re
import pytest
from pathlib import Path
from docx import Document
from src.capabilities import TaskSpecification
from src.document_ir import DocumentIR, SectionIR, TableIR, CitationIR
from src.research_pipeline import clean_wikitext_markup, synthesize_evidence_document
from src.docx_creator import latex_to_omml_xml, create_word_document
from src.ieee_latex_creator import export_to_ieee_latex

def test_defect_1_latex_title_author_no_double_linebreak(tmp_path):
    """Defect 1: LaTeX title/author formatting must not cause 'There's no line here to end'."""
    tex_path = tmp_path / "test_paper.tex"
    doc_dict = {
        "title": "Test Subduction Paper",
        "subtitle": "Kinematic Analysis",
        "abstract": "Test abstract content.",
        "keywords": ["Geodynamics", "Subduction"],
        "author": "[Author Name]",
        "affiliation": "[Author Institution]",
        "location": "[Location]",
        "email": "[author@domain.edu]",
        "sections": [
            {"heading": "I. INTRODUCTION", "level": 1, "paragraphs": ["Introduction paragraph."]}
        ]
    }
    export_to_ieee_latex(doc_dict=doc_dict, output_path=tex_path)
    tex_content = tex_path.read_text(encoding="utf-8")

    # In LaTeX source, consecutive line breaks like \\ \\ or \\\\ cause 'There's no line here to end'
    header = tex_content.split(r"\maketitle")[0]
    assert r"\\ \\" not in header
    assert r"\\\\" not in header
    assert r"\IEEEauthorblockN{\textbf{[Author Name]}}\\" in header
    assert r"\textit{[Author Institution]}, [Location]\\Email: [author@domain.edu]" in header

def test_defect_2_docx_omml_equation_rendering(tmp_path):
    """Defect 2: Equations must be converted to Office MathML (OMML) rather than raw plain text LaTeX."""
    latex_expr = r"v_s = \frac{\Delta \rho}{\rho_m} \cdot \frac{g L^2}{\mu}"
    omml_xml = latex_to_omml_xml(latex_expr)
    assert "<m:oMathPara" in omml_xml
    assert "<m:oMath" in omml_xml
    assert "<m:f>" in omml_xml  # fraction
    assert "<m:num>" in omml_xml
    assert "<m:den>" in omml_xml
    assert "\\mathrm" not in omml_xml

    docx_path = tmp_path / "test_math.docx"
    create_word_document(
        title="Math Test",
        sections=[{
            "heading": "Governing Equations",
            "paragraphs": ["Mathematical formulation below:"],
            "equations": [{"num": 1, "formulas": [latex_expr]}]
        }],
        output_path=docx_path,
        abstract="Math abstract",
        keywords=["Math", "Physics"],
        is_research=True,
        doc_format="ieee"
    )
    doc = Document(str(docx_path))
    xml_str = "".join([p._p.xml for p in doc.paragraphs])
    assert "oMath" in xml_str
    # Verify raw LaTeX commands are not present as plain run text
    assert r"\mathrm" not in xml_str
    assert r"\frac" not in xml_str

def test_defect_3_wikitext_markup_sanitization():
    """Defect 3: Wikitext artifacts like '...|350 px]]' must be completely stripped."""
    dirty_text = "The trench system extends [[File:Hellenic_arc.jpg|thumb|over the Hellenic subduction zone|350 px]] into the Mediterranean."
    cleaned = clean_wikitext_markup(dirty_text)
    assert "350 px]]" not in cleaned
    assert "File:" not in cleaned
    assert "The trench system extends" in cleaned

    orphan_fragment = "Volcanic activity over the Hellenic subduction zone|350 px]] occurred during Miocene."
    cleaned_orphan = clean_wikitext_markup(orphan_fragment)
    assert "350 px]]" not in cleaned_orphan
    assert "350 px" not in cleaned_orphan

def test_defect_4_no_ellipsis_truncation():
    """Defect 4: Comparison table and synthesis text must not truncate findings with ellipsis (...)"""
    task = TaskSpecification(
        action="create_document",
        subject="Subduction Zone Geodynamics",
        deliverable_type="research_paper",
        format="both",
        page_count=6
    )
    doc_ir = synthesize_evidence_document(task)
    for sec in doc_ir.sections:
        for tbl in sec.tables:
            for row in tbl.rows:
                for cell in row:
                    assert not cell.strip().endswith("..."), f"Cell truncated with ellipsis: {cell}"

def test_defect_5_docx_includes_abstract_and_keywords(tmp_path):
    """Defect 5: DOCX must include Abstract and Index Terms matching IEEE structure."""
    docx_path = tmp_path / "test_structure.docx"
    create_word_document(
        title="IEEE Structure Test",
        sections=[{"heading": "Introduction", "paragraphs": ["Body text"]}],
        output_path=docx_path,
        abstract="This is the explicit abstract test.",
        keywords=["Keywords Test 1", "Keywords Test 2"],
        is_research=True,
        doc_format="ieee"
    )
    doc = Document(str(docx_path))
    full_text = "\n".join([p.text for p in doc.paragraphs])
    assert "Abstract—" in full_text
    assert "This is the explicit abstract test." in full_text
    assert "Index Terms—" in full_text
    assert "Keywords Test 1, Keywords Test 2" in full_text

def test_defect_6_honest_methodology_framing():
    """Defect 6: Pipeline must frame methodology honestly as literature synthesis and document queries."""
    task = TaskSpecification(
        action="create_document",
        subject="Subduction Zone Geodynamics",
        deliverable_type="research_paper",
        format="both",
        page_count=6
    )
    doc_ir = synthesize_evidence_document(task)
    # Ensure title does not falsely claim "A Systematic Review"
    assert "Systematic Review" not in doc_ir.title
    assert "Literature Synthesis" in doc_ir.title or "Evidence-Grounded" in doc_ir.title
    # Ensure methodology section is present and documents search queries and criteria
    methodology_sections = [s for s in doc_ir.sections if "METHODOLOGY" in s.heading.upper()]
    assert len(methodology_sections) > 0
    meth_content = " ".join(methodology_sections[0].paragraphs)
    assert "arXiv" in meth_content
    assert "MediaWiki" in meth_content
    assert "Inclusion and screening criteria" in meth_content

def test_defect_7_citation_integrity_and_attribution():
    """Defect 7: Wikipedia attributions must cite 'Wikipedia contributors' (not fake consortium), HTTPS URLs, valid DOIs."""
    task = TaskSpecification(
        action="create_document",
        subject="Subduction Zone Geodynamics",
        deliverable_type="research_paper",
        format="both",
        page_count=6
    )
    doc_ir = synthesize_evidence_document(task)
    assert len(doc_ir.bibliography) > 0
    for ref in doc_ir.bibliography:
        assert "Encyclopedia Research Consortium" not in ref.authors
        if "Wikipedia" in ref.title or ref.source_type == "wikipedia":
            assert "Wikipedia contributors" in ref.authors
            assert ref.source_url.startswith("https://")
        if ref.doi:
            assert not ref.doi.startswith("doi:")
            assert ref.doi.startswith("10.")
