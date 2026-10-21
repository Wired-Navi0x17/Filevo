"""Tests for comprehensive Markdown parser and DOCX renderer."""

import tempfile
from pathlib import Path
from docx import Document
from docx.oxml.ns import qn

from src.docx_creator import (
    add_math_equation_block,
    create_word_document,
    parse_document_request,
)
from src.markdown_parser import (
    add_inline_markdown,
    markdown_to_docx,
    parse_markdown_into_doc,
    parse_markdown_into_sections,
    render_alert_box,
    render_markdown_table,
)
from src.extractors import extract_document


def test_inline_markdown_formatting():
    """Verify inline tokens: bold, italic, strikethrough, code, highlight, sup/sub, math, link."""
    doc = Document()
    p = doc.add_paragraph()
    text = (
        "Here is **bold**, *italic*, ***bold italic***, ~~struck~~, `code`, "
        "==highlight==, <sup>sup</sup>, <sub>sub</sub>, $E = mc^2$, and [Google](https://google.com)."
    )
    add_inline_markdown(p, text)

    assert len(p.runs) >= 10
    # Bold run
    bold_runs = [r for r in p.runs if r.font.bold and not r.font.italic]
    assert any("bold" in r.text for r in bold_runs)

    # Italic run
    italic_runs = [r for r in p.runs if r.font.italic and not r.font.bold]
    assert any("italic" in r.text for r in italic_runs)

    # Bold Italic run
    bi_runs = [r for r in p.runs if r.font.bold and r.font.italic]
    assert any("bold italic" in r.text for r in bi_runs)

    # Strikethrough
    strike_runs = [r for r in p.runs if r.font.strike]
    assert any("struck" in r.text for r in strike_runs)

    # Monospace code span
    code_runs = [r for r in p.runs if r.font.name == "Courier New"]
    assert any("code" in r.text for r in code_runs)

    # Subscript & Superscript
    sub_runs = [r for r in p.runs if r.font.subscript]
    assert any("sub" in r.text for r in sub_runs)

    sup_runs = [r for r in p.runs if r.font.superscript]
    assert any("sup" in r.text for r in sup_runs)

    # Link run
    link_runs = [r for r in p.runs if r.font.underline]
    assert any("Google" in r.text for r in link_runs)


def test_alert_callouts_rendering():
    """Verify GitHub / Obsidian styled alert callout boxes."""
    doc = Document()
    for alert_type in ("NOTE", "TIP", "IMPORTANT", "WARNING", "CAUTION"):
        table = render_alert_box(
            doc,
            alert_type=alert_type,
            title=f"Custom {alert_type}",
            lines=[f"This is an urgent {alert_type} line.", "Second detail paragraph."],
        )
        assert table is not None
        assert len(table.rows) == 1
        cell = table.rows[0].cells[0]
        full_text = " ".join(p.text for p in cell.paragraphs)
        assert alert_type in full_text
        assert "Second detail" in full_text


def test_gfm_table_rendering_and_alignments():
    """Verify GFM Markdown table parsing and DOCX rendering with alignments."""
    doc = Document()
    headers = ["Task", "Status", "Priority"]
    rows = [
        ["Policy Gradient", "Completed", "High"],
        ["Softmax Tuning", "In Progress", "Medium"],
        ["LaTeX Transpiler", "Verified", "Critical"],
    ]
    alignments = ["left", "center", "right"]

    tbl = render_markdown_table(doc, headers, rows, alignments)
    assert tbl is not None
    assert len(tbl.rows) == 4  # 1 header + 3 data rows
    assert len(tbl.columns) == 3


def test_math_equation_block_has_no_interior_vertical_borders():
    """Verify mathematical equation table blocks have insideV='none' to avoid vertical divider lines."""
    doc = Document()
    formula_lines = [
        "π_θ(a | x) = exp(θ_a^⊤ x / τ) / ∑_{j=1}^K exp(θ_j^⊤ x / τ)",
        "θ_{t+1} = θ_t + α · R(x, a) · ∇_θ log π_θ(a | x)",
    ]
    table = add_math_equation_block(
        doc,
        formula_lines=formula_lines,
        eq_num="1",
        explanation="Gibbs softmax policy gradient parameterization.",
    )
    assert table is not None
    assert len(table.columns) == 2

    # Inspect OXML table borders
    tblBorders = table._tbl.tblPr.find(qn("w:tblBorders"))
    assert tblBorders is not None
    insideV = tblBorders.find(qn("w:insideV"))
    assert insideV is not None
    assert insideV.get(qn("w:val")) == "none"


def test_parse_markdown_into_sections():
    """Verify full markdown string parsing into Sortwise section dictionaries."""
    sample_md = """# Introduction to Reinforcement Learning
Reinforcement learning agents optimize long-term return $J(\\theta)$.

- [ ] Implement score function
- [x] Verify gradient ascent
- Analytical convergence

> [!NOTE]
> Parameter updates are proportional to feedback rewards $R \\in [-1, +1]$.

## Mathematical Formulations
Gibbs policy distribution:
$$
\\pi_\\theta(a | x) = \\frac{\\exp(\\theta_a^\\top x / \\tau)}{\\sum_j \\exp(\\theta_j^\\top x / \\tau)}
$$

```python
def update_step(theta, grad, lr=0.01):
    return theta + lr * grad
```

| Hyperparameter | Value | Description |
| :--- | :---: | ---: |
| Learning Rate $\\alpha$ | 0.05 | Step size |
| Temperature $\\tau$ | 0.15 | Exploration |
"""
    sections = parse_markdown_into_sections(sample_md)
    assert len(sections) == 2

    sec1 = sections[0]
    assert sec1["heading"] == "Introduction to Reinforcement Learning"
    assert sec1["level"] == 1
    assert any("Reinforcement learning" in p for p in sec1["paragraphs"])
    assert any("☑" in b for b in sec1["bullets"])  # checked task
    assert any("☐" in b for b in sec1["bullets"])  # unchecked task
    assert len(sec1["alerts"]) == 1
    assert sec1["alerts"][0]["type"] == "NOTE"

    sec2 = sections[1]
    assert sec2["heading"] == "Mathematical Formulations"
    assert sec2["level"] == 2
    assert len(sec2["equations"]) == 1
    assert sec2["code"] is not None
    assert "def update_step" in sec2["code"]
    assert sec2["table"] is not None
    assert sec2["table"]["headers"] == ["Hyperparameter", "Value", "Description"]
    assert sec2["table"]["alignments"] == ["left", "center", "right"]


def test_markdown_to_docx_file_export():
    """Verify markdown_to_docx outputs a valid A4 document with all elements."""
    with tempfile.TemporaryDirectory() as tmp_dir:
        out_docx = Path(tmp_dir) / "full_markdown_test.docx"
        md_text = """# Executive System Specification

> [!TIP]
> Always verify mathematical calibration before production deployments.

## Architecture & Data Flow
Sortwise uses calibrated entropy metrics:

- High confidence: direct automated dispatch
- High entropy: active learning query

| Module | Metric | Threshold |
| :--- | :---: | ---: |
| Classifier | $H_{norm}$ | 0.35 |
| Calibrator | $D_{JS}$ | 0.12 |
"""
        doc = markdown_to_docx(md_text, output_path=out_docx, title="Sortwise Architecture Report")
        assert out_docx.exists()
        assert out_docx.stat().st_size > 5000

        edoc = extract_document(str(out_docx))
        assert "Sortwise Architecture Report" in edoc.title or "Executive System Specification" in edoc.clean_text
        assert "TIP" in edoc.clean_text
        assert "High confidence" in edoc.clean_text


def test_parse_document_request_with_attached_markdown_file():
    """Verify parse_document_request directly parses attached .md files."""
    with tempfile.TemporaryDirectory() as tmp_dir:
        md_file = Path(tmp_dir) / "research_notes.md"
        md_file.write_text(
            "# Autonomous Policy Synthesis\n\n"
            "This document outlines continuous gradient ascent.\n\n"
            "## Methodology\n"
            "Active learning queries resolve ambiguity.\n",
            encoding="utf-8",
        )

        parsed = parse_document_request(
            "create a word document from research_notes.md",
            base_dir=Path(tmp_dir),
            attachment=md_file,
        )
        assert parsed["title"] == "Autonomous Policy Synthesis"
        assert len(parsed["sections"]) == 2
        assert parsed["sections"][0]["heading"] == "Autonomous Policy Synthesis"
        assert parsed["sections"][1]["heading"] == "Methodology"
