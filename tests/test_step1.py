"""Step 1 acceptance tests.

Generates a Calculus PDF, an Algorithms DOCX and a Cyber Security PPTX on the
fly, then checks that:

* all three extractors produce a clean :class:`ExtractedDoc`, and
* the calibrated classifier routes each document to the right subject with a
  well-formed confidence / entropy / status decision.
"""

from __future__ import annotations

import math
import sys
from pathlib import Path

import pytest

# Make ``src`` importable no matter where pytest is invoked from.
PROJECT_ROOT = Path(__file__).resolve().parents[1]
if str(PROJECT_ROOT) not in sys.path:
    sys.path.insert(0, str(PROJECT_ROOT))

import fitz  # PyMuPDF
import docx  # python-docx
from pptx import Presentation

from src.extractors import (
    DocumentExtractionError,
    ExtractedDoc,
    extract_document,
)
from src.taxonomy import get_subject_taxonomy
from src.calibrator import (
    CONFIDENCE_THRESHOLD,
    ENTROPY_THRESHOLD,
    STATUS_AMBIGUOUS,
    STATUS_CONFIDENT,
    CalibratedClassifier,
    CalibratedDecision,
    brier_score,
    shannon_entropy,
    tempered_softmax,
)

CALCULUS_LABEL = "CS2806 - Calculus"
ALGORITHMS_LABEL = "CS2000 - Design and Analysis of Algorithms"
CYBER_LABEL = "CS2405 - Cyber Security"


# --------------------------------------------------------------------------- #
# Sample document generation
# --------------------------------------------------------------------------- #
CALCULUS_BODY = """\
1. Limits and Continuity
The limit of a function describes its behaviour as the input approaches a
value. A function is continuous when the limit equals the function value.

2. Differential Calculus
The derivative measures the instantaneous rate of change of a function.
Differentiation rules include the power rule, the product rule, the quotient
rule and the chain rule. Implicit differentiation is used when y is not given
explicitly. Critical points, maxima and minima, concavity and inflection
points are found with the first and second derivatives.

3. Applications of Derivatives
Related rates, optimisation and the mean value theorem rely on derivatives.
L'Hopital's rule evaluates indeterminate limits.

4. Integral Calculus
An antiderivative reverses differentiation. The Riemann sum defines the
definite integral as the area under a curve. The fundamental theorem of
calculus links derivatives and integrals. Techniques include substitution and
integration by parts, and applications include volumes of revolution.

5. Sequences and Series
Sequences and infinite series are tested for convergence and divergence using
the ratio and comparison tests. Taylor series and Maclaurin series approximate
functions.
"""

ALGORITHMS_PARAGRAPHS = [
    ("Sorting Algorithms and Complexity", 1),
    ("Merge sort and quick sort divide the input and combine sorted parts, "
     "achieving O(n log n) average time complexity, while insertion sort and "
     "bubble sort are quadratic. Heap sort builds a binary heap and repeatedly "
     "extracts the maximum. Counting sort and radix sort run in linear time for "
     "bounded integer keys.", None),
    ("Asymptotic Analysis", 1),
    ("Big-O notation describes the asymptotic upper bound of time and space "
     "complexity. Recurrence relations such as T(n) = 2T(n/2) + n are solved "
     "with the master theorem. Amortized analysis bounds the average cost of "
     "operations on dynamic arrays and heaps.", None),
    ("Algorithm Design Paradigms", 1),
    ("Divide and conquer splits a problem into subproblems. Greedy algorithms "
     "make locally optimal choices, as in Kruskal and Prim for the minimum "
     "spanning tree. Dynamic programming solves overlapping subproblems, for "
     "example shortest paths with Bellman-Ford and Floyd-Warshall.", None),
    ("Graphs and Complexity Classes", 1),
    ("Graph algorithms include breadth first search, depth first search, "
     "Dijkstra's shortest path and topological sort. Complexity theory studies "
     "polynomial reductions, NP-completeness and approximation algorithms.",
     None),
]

CYBER_SLIDES = [
    ("Cyber Security: Cryptography and Network Attacks",
     "Symmetric and asymmetric encryption, hashing, and common network attacks."),
    ("Cryptography Fundamentals",
     "Symmetric key encryption uses AES and DES. Asymmetric public key "
     "cryptography uses RSA and elliptic curves. Hash functions such as SHA "
     "provide integrity, while digital signatures and HMAC provide message "
     "authentication. Diffie-Hellman performs key exchange."),
    ("Network Attacks",
     "Malware includes viruses, worms, trojans and ransomware. Phishing and "
     "social engineering target users. Denial of service and DDoS exhaust "
     "resources. Man-in-the-middle, SQL injection, cross-site scripting and "
     "buffer overflow exploit vulnerable systems."),
    ("Defences",
     "Firewalls, intrusion detection, TLS and public key infrastructure, access "
     "control, threat modelling and penetration testing defend networks."),
]


def _write_pdf(path: Path) -> None:
    doc = fitz.open()
    page = doc.new_page(width=595, height=842)
    page.insert_text((72, 60), "CS2806 - Calculus", fontsize=20)
    page.insert_text((72, 88), "Differential and Integral Calculus", fontsize=14)
    page.insert_textbox(fitz.Rect(72, 110, 523, 800), CALCULUS_BODY, fontsize=11)
    doc.set_metadata({"title": "Calculus Lecture Notes"})
    doc.save(str(path))
    doc.close()


def _write_docx(path: Path) -> None:
    document = docx.Document()
    document.core_properties.title = "Algorithms Course Notes"
    for text, level in ALGORITHMS_PARAGRAPHS:
        if level is None:
            document.add_paragraph(text)
        else:
            document.add_heading(text, level=level)
    document.save(str(path))


def _write_pptx(path: Path) -> None:
    presentation = Presentation()
    title_layout = presentation.slide_layouts[0]
    content_layout = presentation.slide_layouts[1]
    for index, (title, body) in enumerate(CYBER_SLIDES):
        layout = title_layout if index == 0 else content_layout
        slide = presentation.slides.add_slide(layout)
        slide.shapes.title.text = title
        slide.placeholders[1].text = body
    presentation.core_properties.title = "Cyber Security Slides"
    presentation.save(str(path))


@pytest.fixture(scope="module")
def sample_docs(tmp_path_factory: pytest.TempPathFactory) -> dict[str, Path]:
    base = tmp_path_factory.mktemp("step1_samples")
    pdf_path = base / "calculus.pdf"
    docx_path = base / "algorithms.docx"
    pptx_path = base / "cyber_security.pptx"
    _write_pdf(pdf_path)
    _write_docx(docx_path)
    _write_pptx(pptx_path)
    return {"pdf": pdf_path, "docx": docx_path, "pptx": pptx_path}


# --------------------------------------------------------------------------- #
# Step 1a -- extractors
# --------------------------------------------------------------------------- #
def test_extract_pdf(sample_docs: dict[str, Path]) -> None:
    doc = extract_document(str(sample_docs["pdf"]))

    assert isinstance(doc, ExtractedDoc)
    assert doc.file_type == "pdf"
    assert doc.title == "Calculus Lecture Notes"
    assert doc.word_count > 50
    assert doc.word_count == len(doc.clean_text.split())
    lowered = doc.clean_text.lower()
    assert "derivative" in lowered
    assert "integral" in lowered
    # Layout heading detection should pick up the large title text.
    assert any("calculus" in heading.lower() for heading in doc.headings)


def test_extract_docx(sample_docs: dict[str, Path]) -> None:
    doc = extract_document(str(sample_docs["docx"]))

    assert doc.file_type == "docx"
    assert doc.title == "Algorithms Course Notes"
    assert doc.word_count > 50
    assert doc.headings, "DOCX headings should not be empty"
    assert any("sorting" in heading.lower() for heading in doc.headings)
    assert "dynamic programming" in doc.clean_text.lower()


def test_extract_pptx(sample_docs: dict[str, Path]) -> None:
    doc = extract_document(str(sample_docs["pptx"]))

    assert doc.file_type == "pptx"
    assert doc.title == "Cyber Security Slides"
    assert doc.word_count > 30
    assert doc.headings, "PPTX slide titles should not be empty"
    assert any("cryptography" in heading.lower() for heading in doc.headings)
    assert "rsa" in doc.clean_text.lower()


def test_extract_missing_file_raises(tmp_path: Path) -> None:
    with pytest.raises(DocumentExtractionError):
        extract_document(str(tmp_path / "does_not_exist.pdf"))


def test_extract_unsupported_type_raises(tmp_path: Path) -> None:
    bogus = tmp_path / "notes.txt"
    bogus.write_text("just some text", encoding="utf-8")
    with pytest.raises(DocumentExtractionError):
        extract_document(str(bogus))


# --------------------------------------------------------------------------- #
# Step 1b -- taxonomy
# --------------------------------------------------------------------------- #
def test_taxonomy_contains_all_eight_subjects() -> None:
    taxonomy = get_subject_taxonomy()
    assert len(taxonomy) == 8
    codes = {entry["code"] for entry in taxonomy.values()}
    assert codes == {
        "CS2000",
        "CS2227",
        "CS2235",
        "CS2403",
        "CS2404",
        "CS2405",
        "CS2500",
        "CS2806",
    }
    assert CALCULUS_LABEL in taxonomy
    for label, entry in taxonomy.items():
        assert entry["keywords"], f"{label} has no keywords"
        assert entry["syllabus"].strip(), f"{label} has no syllabus"


def test_taxonomy_is_a_copy() -> None:
    first = get_subject_taxonomy()
    first[CALCULUS_LABEL]["keywords"].append("mutation")
    second = get_subject_taxonomy()
    assert "mutation" not in second[CALCULUS_LABEL]["keywords"]


# --------------------------------------------------------------------------- #
# Step 1c -- calibrator
# --------------------------------------------------------------------------- #
@pytest.fixture(scope="module")
def classifier() -> CalibratedClassifier:
    return CalibratedClassifier()


EXPECTED = {"pdf": CALCULUS_LABEL, "docx": ALGORITHMS_LABEL, "pptx": CYBER_LABEL}


@pytest.mark.parametrize("kind", ["pdf", "docx", "pptx"])
def test_calibrator_routes_documents(
    sample_docs: dict[str, Path],
    classifier: CalibratedClassifier,
    kind: str,
) -> None:
    doc = extract_document(str(sample_docs[kind]))
    decision = classifier.classify_document(doc)

    assert isinstance(decision, CalibratedDecision)
    assert decision.subject == EXPECTED[kind], (
        f"{kind}: expected {EXPECTED[kind]}, got {decision.subject} "
        f"(confidence={decision.confidence:.3f}, entropy={decision.entropy:.3f})"
    )
    # The correct subject must actually be the best-scoring one.
    assert decision.ranking[0][0] == decision.subject


def test_decision_is_well_formed(
    sample_docs: dict[str, Path], classifier: CalibratedClassifier
) -> None:
    doc = extract_document(str(sample_docs["pdf"]))
    decision = classifier.classify_document(doc)

    n_subjects = len(get_subject_taxonomy())
    assert 0.0 <= decision.confidence <= 1.0
    assert 0.0 <= decision.entropy <= math.log2(n_subjects) + 1e-9
    assert decision.status in {STATUS_CONFIDENT, STATUS_AMBIGUOUS}
    assert len(decision.ranking) == n_subjects

    # Ranking is sorted descending and is a proper probability distribution.
    scores = [score for _, score in decision.ranking]
    assert scores == sorted(scores, reverse=True)
    assert math.isclose(sum(scores), 1.0, rel_tol=1e-6)
    assert math.isclose(decision.confidence, scores[0], rel_tol=1e-9)


def test_status_matches_triage_rule(
    sample_docs: dict[str, Path], classifier: CalibratedClassifier
) -> None:
    for kind in EXPECTED:
        doc = extract_document(str(sample_docs[kind]))
        decision = classifier.classify_document(doc)
        confident = (
            decision.confidence >= CONFIDENCE_THRESHOLD
            and decision.entropy < ENTROPY_THRESHOLD
        )
        expected_status = STATUS_CONFIDENT if confident else STATUS_AMBIGUOUS
        assert decision.status == expected_status


def test_clear_documents_are_confident(
    sample_docs: dict[str, Path], classifier: CalibratedClassifier
) -> None:
    """Purpose-built, on-topic documents should clear the triage gate."""
    for kind in EXPECTED:
        doc = extract_document(str(sample_docs[kind]))
        decision = classifier.classify_document(doc)
        assert decision.status == STATUS_CONFIDENT, (
            f"{kind} was {decision.status}: confidence={decision.confidence:.3f}, "
            f"entropy={decision.entropy:.3f}"
        )


def test_low_temperature_sharpens_distribution() -> None:
    similarities = [0.5, 0.4, 0.2, 0.1]
    cold = tempered_softmax(similarities, 0.15)
    warm = tempered_softmax(similarities, 1.0)
    assert cold.max() > warm.max()
    assert math.isclose(float(cold.sum()), 1.0, rel_tol=1e-9)


def test_shannon_entropy_bounds() -> None:
    uniform = [0.25, 0.25, 0.25, 0.25]
    assert math.isclose(shannon_entropy(uniform), 2.0, rel_tol=1e-9)
    assert math.isclose(shannon_entropy([1.0, 0.0, 0.0, 0.0]), 0.0, abs_tol=1e-9)


def test_brier_score() -> None:
    # Perfect, fully confident correct prediction.
    assert brier_score([1.0, 0.0, 0.0], 0) == pytest.approx(0.0)
    # Confident but wrong.
    assert brier_score([1.0, 0.0, 0.0], 1) == pytest.approx(2.0)
    # Uniform guess over 4 classes.
    assert brier_score([0.25, 0.25, 0.25, 0.25], 0) == pytest.approx(0.75)
    with pytest.raises(ValueError):
        brier_score([0.5, 0.5], 5)
    with pytest.raises(ValueError):
        brier_score([], 0)
